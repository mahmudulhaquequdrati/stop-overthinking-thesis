"""Populate cached automatic Word navigation from verified rendered pagination."""
import json,re,sys
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt
from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER

base=Path(__file__).resolve().parents[1]
docpath=base/'thesis.docx'
r=PdfReader(sys.argv[1]);texts=[p.extract_text() for p in r.pages]
norm=lambda x:re.sub('[^a-z0-9]','',x.lower())
start=next(i for i,t in enumerate(texts) if any(line.strip()=='Chapter 1: Introduction' for line in t.splitlines()))
body=[norm(t) for t in texts[start:]]
def pageof(s):
 n=norm(s)
 for i,t in enumerate(body):
  if s=='References':
   if any(norm(line)==n for line in texts[start+i].splitlines()):return i+1
  elif n in t:return i+1
 raise RuntimeError('Missing rendered navigation target '+s)
blocks=json.loads((base/'source/thesis-content.json').read_text())
caps=json.loads((base/'qa/captions.json').read_text())
nav=[]
for k,t in blocks:
 if k in ['chapter','heading']:nav.append((t,pageof(t),k=='chapter'))
nav += [(t,pageof(t),True) for t in ['References','Appendix A: Evidence and Supplementary Data','Appendix B: Reproduction Procedure','Appendix C: Further Checks and Authorship Details']]
figs=[]
for key,n in caps['figure_numbers'].items():
 cap=caps['figures'][key];page=pageof('Figure '+str(n)+'. '+cap[:45]);figs.append((f'Figure {n}. '+cap.split('. ')[0]+'.',page,False))
tables=[]
for key,n in caps['table_numbers'].items():
 # The source captions follow the automatic number as normal text.
 d=Document(docpath)
 needle='Table '+str(n)+'. '
 for p in d.paragraphs:
  if p.style is not None and p.style.name=='Caption' and p.text.startswith(needle):
   tables.append((p.text,pageof(p.text[:90]),False));break
d=Document(docpath)
def runxml(p,tag,text=None,**attrs):
 r=OxmlElement('w:r');e=OxmlElement('w:'+tag)
 for k,v in attrs.items():e.set(qn('w:'+k),v)
 if text is not None:e.text=text
 r.append(e);p._p.append(r)
for title,nexttitle,code,entries in [('Table of Contents','List of Figures','TOC \\o "1-2" \\h \\z \\u',nav),('List of Figures','List of Tables','TOC \\h \\z \\c "Figure"',figs),('List of Tables','List of Abbreviations','TOC \\h \\z \\c "Table"',tables)]:
 for p in list(d.paragraphs):
  if p.text!=title:continue
  nxt=next(q for q in d.paragraphs if q.text==nexttitle)
  anchor=nxt._p
  nxt.paragraph_format.page_break_before=True
  el=p._p.getnext()
  while el is not None and el is not anchor:
   following=el.getnext();el.getparent().remove(el);el=following
  generated=[]
  for title,page,bold in entries:
   np=d.add_paragraph();np.paragraph_format.line_spacing=1.05;np.paragraph_format.space_after=Pt(5);np.paragraph_format.space_before=Pt(0);np.paragraph_format.keep_with_next=False;np.paragraph_format.keep_together=True
   np.paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
   if not bold and entries is nav:np.paragraph_format.left_indent=Inches(.15)
   rr=np.add_run(title+'\t'+str(page));rr.font.size=Pt(11);rr.bold=bold
   anchor.addprevious(np._p);generated.append(np)
  first,last=generated[0],generated[-1]
  # The enclosing field remains a genuine automatic TOC/figure/table field.
  prefix=[]
  for typ,value in [('fldChar','begin'),('instrText',code),('fldChar','separate')]:
   rr=OxmlElement('w:r');e=OxmlElement('w:'+typ)
   if typ=='fldChar':e.set(qn('w:fldCharType'),value)
   else:e.text=value;e.set(qn('xml:space'),'preserve')
   rr.append(e);prefix.append(rr)
  for ix,rr in enumerate(prefix):first._p.insert(ix+1,rr)
  runxml(last,'fldChar',fldCharType='end')
  break
d.save(docpath)
(base/'qa/navigation-pages.json').write_text(json.dumps({'main_start_physical':start+1,'contents':nav,'figures':figs,'tables':tables},indent=2))
print('Cached navigation:',len(nav),'contents entries,',len(figs),'figures,',len(tables),'tables')
